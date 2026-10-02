"""
ReportAgent — synthèse finale.

Transforme les scores (scoring.py) et les Finding persistés en un rapport
structuré : résumé exécutif (non-technique), annexe technique, plan de
remédiation trié par priorité, et export Excel des findings bruts.

Le PDF bilingue FR/AR n'est pas encore implémenté — mise en page et typographie
arabe (RTL, police dédiée) sont un chantier à part, volontairement reporté
plutôt que livré à la hâte.
"""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

CATEGORY_LABELS = {
    "known_vulns": "Vulnérabilités connues",
    "exposed_services": "Services exposés",
    "tls": "TLS / Certificat",
    "headers": "Headers de sécurité",
    "email": "Sécurité email",
}

SEVERITY_LABELS = {1: "Critique", 2: "Élevée", 3: "Moyenne", 4: "Faible", 5: "Informatif"}


def _score_level(score: float) -> str:
    if score >= 80:
        return "bon"
    if score >= 50:
        return "moyen"
    return "préoccupant"


def generate_executive_summary(
    domain: str,
    global_score: float,
    category_scores: dict[str, float],
    findings: list[dict[str, Any]],
) -> str:
    """Résumé en langage simple, pour un lecteur non technique."""
    confirmed_critical = [
        f for f in findings
        if f.get("confidence") == "confirmed" and f.get("business_priority") in (1, 2)
    ]
    worst_category = min(category_scores, key=category_scores.get) if category_scores else None

    lines = [
        f"Audit de sécurité de {domain}",
        "",
        f"Score global : {global_score}/100 ({_score_level(global_score)}).",
        "",
    ]

    if confirmed_critical:
        lines.append(
            f"{len(confirmed_critical)} problème(s) critique(s) ou élevé(s) ont été confirmés "
            "et nécessitent une action rapide."
        )
    else:
        lines.append("Aucun problème critique confirmé n'a été détecté lors de cet audit.")

    if worst_category and category_scores[worst_category] < 70:
        lines.append(
            f"Le point le plus faible concerne la catégorie "
            f"« {CATEGORY_LABELS.get(worst_category, worst_category)} » "
            f"({category_scores[worst_category]}/100)."
        )

    needs_review = sum(1 for f in findings if f.get("confidence") == "needs_manual_review")
    if needs_review:
        lines.append(
            f"{needs_review} résultat(s) supplémentaire(s) nécessitent une vérification manuelle : "
            "un scan automatique ne peut pas confirmer à 100% qu'une faille est réellement exploitable."
        )

    return "\n".join(lines)


def generate_technical_annex(
    category_scores: dict[str, float],
    findings: list[dict[str, Any]],
    scan_results: dict[str, Any] | None,
    recon_results: dict[str, Any] | None,
) -> dict[str, Any]:
    """Détail technique complet, pour l'équipe IT — structuré, pas prosifié."""
    scan_results = scan_results or {}
    recon_results = recon_results or {}

    return {
        "category_scores": category_scores,
        "findings_by_category": {
            category: [f for f in findings if f.get("category") == category]
            for category in CATEGORY_LABELS
        },
        "open_ports": scan_results.get("ports", {}).get("open_ports", []),
        "tls_details": scan_results.get("tls", {}),
        "subdomains_found": recon_results.get("subdomain_count", 0),
        "dangling_dns_candidates": recon_results.get("dangling_dns", []),
        "email_security": recon_results.get("email_security", {}),
    }


def generate_remediation_plan(findings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Plan trié par priorité business (1 = le plus urgent), pas par ordre d'apparition."""
    sorted_findings = sorted(
        (f for f in findings if f.get("remediation")),
        key=lambda f: (f.get("business_priority") or 5, f.get("confidence") != "confirmed"),
    )

    return [
        {
            "priority": f.get("business_priority"),
            "priority_label": SEVERITY_LABELS.get(f.get("business_priority"), "Non classé"),
            "title": f.get("title"),
            "remediation": f.get("remediation"),
            "confidence": f.get("confidence"),
            "cve_id": f.get("cve_id"),
        }
        for f in sorted_findings
    ]


def export_findings_excel(
    domain: str,
    findings: list[dict[str, Any]],
    category_scores: dict[str, float],
) -> bytes:
    """Génère un classeur Excel des findings bruts, en mémoire (pas de fichier temporaire)."""
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Résumé"
    summary_sheet.append(["Catégorie", "Score / 100"])
    for cell in summary_sheet[1]:
        cell.font = Font(bold=True)
    for category, score in category_scores.items():
        summary_sheet.append([CATEGORY_LABELS.get(category, category), score])

    findings_sheet = workbook.create_sheet("Findings")
    headers = ["Catégorie", "Titre", "Sévérité CVSS", "Priorité", "Confiance", "CVE", "Remédiation"]
    findings_sheet.append(headers)
    header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    for cell in findings_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for finding in findings:
        findings_sheet.append([
            CATEGORY_LABELS.get(finding.get("category"), finding.get("category")),
            finding.get("title"),
            finding.get("severity"),
            finding.get("business_priority"),
            finding.get("confidence"),
            finding.get("cve_id"),
            finding.get("remediation"),
        ])

    for sheet in (summary_sheet, findings_sheet):
        for column_cells in sheet.columns:
            length = max((len(str(c.value)) for c in column_cells if c.value), default=10)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 60)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_score_history_entries(audit_id: str, domain: str, category_scores: dict[str, float]) -> list[dict[str, Any]]:
    """Prépare les lignes ScoreHistory à insérer — un snapshot par catégorie,
    horodaté, pour que le score reste comparable même si la formule de calcul
    change plus tard (voir le commentaire dans models.py).
    """
    now = datetime.now(timezone.utc)
    return [
        {"audit_id": audit_id, "domain": domain, "category": category, "score": score, "recorded_at": now}
        for category, score in category_scores.items()
    ]


def generate_report(
    audit_id: str,
    domain: str,
    findings: list[dict[str, Any]],
    scan_results: dict[str, Any] | None,
    recon_results: dict[str, Any] | None,
) -> dict[str, Any]:
    """Point d'entrée du ReportAgent : assemble toutes les briques en un seul résultat."""
    from app.agents.report.scoring import compute_category_scores, compute_global_score

    category_scores = compute_category_scores(findings, scan_results, recon_results)
    global_score = compute_global_score(category_scores)

    return {
        "global_score": global_score,
        "category_scores": category_scores,
        "executive_summary": generate_executive_summary(domain, global_score, category_scores, findings),
        "technical_annex": generate_technical_annex(category_scores, findings, scan_results, recon_results),
        "remediation_plan": generate_remediation_plan(findings),
        "score_history_entries": build_score_history_entries(audit_id, domain, category_scores),
        # bytes Excel non inclus ici : généré à la demande par l'endpoint,
        # pas stocké en base (trop volumineux pour une colonne JSON).
    }